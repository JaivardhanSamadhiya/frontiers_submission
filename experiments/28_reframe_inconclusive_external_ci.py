#!/usr/bin/env python3
"""Reframe the frozen cross-source result as statistically inconclusive."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "submission/PrecisionPhage_Frontiers_Original_Research_frozen_external_validated.docx"
OUTPUT = ROOT / "submission/PrecisionPhage_Frontiers_Original_Research_cross_source_uncertainty_reframed.docx"


def replace_paragraph(doc: Document, starts_with: str, replacement: str) -> None:
    matches = [p for p in doc.paragraphs if p.text.startswith(starts_with)]
    if len(matches) != 1:
        raise AssertionError(
            f"expected one paragraph starting {starts_with!r}; found {len(matches)}"
        )
    paragraph = matches[0]
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(replacement)


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    replace_paragraph(
        doc,
        "Leakage-controlled phage–host interaction prediction reveals",
        "Leakage-controlled phage–host interaction prediction reveals uncertainty in cross-source transfer",
    )
    replace_paragraph(
        doc,
        "Computational prediction of bacteriophage–host interactions",
        "Computational prediction of bacteriophage–host interactions may prioritize experimental host-range testing, but related genomes and source structure can inflate evaluation. We present PrecisionPhage, a reproducible leakage-controlled framework using 8,849 experimentally assayed pairs from three study labels. Within the 1,947-pair sequence-covered NCBI_HR subset (1,488 positive and 459 negative), pooled fixed-tree AUROC decreased from 0.959 for eligible unseen host species to 0.785 across five dual cold-start blocks. We then locked the feature pipeline and model before retraining-free evaluation on all 1,053 sequence-covered StaphStudy pairs (333 positive and 720 negative). The full representation yielded AUROC 0.637 (two-way sequence-cluster bootstrap 95% CI 0.498–0.739) and AUPRC 0.391 (0.201–0.591) against 31.6% positive prevalence. Because the primary AUROC interval includes 0.5, this test does not resolve whether discrimination exceeds chance. A prespecified sequence-only sensitivity model excluding four source-supplied features gave a similar point estimate (AUROC 0.639), whereas those four features alone gave AUROC 0.444. The full model was poorly calibrated (expected calibration error 0.499). On a secondary 805-pair subset beyond the prespecified phage and host similarity threshold, AUROC was 0.644. Nested-threshold predictions within NCBI_HR yielded a 176-phage retrospective set cover with 89.6% observed at-least-one coverage. An unfitted deterministic parameter grid was retained only as an assumption-propagation illustration. These results establish a within-source leakage hierarchy, but cross-source discrimination remains inconclusive; they do not support a generally calibrated predictor, resistance mechanism, or therapeutic efficacy claim.",
    )
    replace_paragraph(
        doc,
        "3.2 Frozen source-held-out evaluation",
        "3.2 Frozen source-held-out evaluation is inconclusive and poorly calibrated",
    )
    replace_paragraph(
        doc,
        "The frozen fixed GBM trained on 1,947 NCBI_HR pairs",
        "The frozen fixed GBM trained on 1,947 NCBI_HR pairs and applied once to 1,053 StaphStudy pairs yielded AUROC 0.637 (two-way sequence-cluster bootstrap 95% CI 0.498–0.739) and AUPRC 0.391 (0.201–0.591), compared with 31.6% positive prevalence (Table 8). The primary AUROC interval includes 0.5; therefore, this evaluation does not provide statistically resolved evidence that the full model discriminates above chance across this source boundary. The sequence-only sensitivity model excluding k3dist, k6dist, GCdiff, and Homology had a similar point estimate (AUROC 0.639; AUPRC 0.392), whereas those four source-supplied features alone transferred poorly (AUROC 0.444; AUPRC 0.270). The sensitivity analyses provide context but do not supersede the prespecified primary result. The full model was not transportably calibrated (Brier score 0.466; ECE 0.499). All 39 test phage identifiers and all exact genome sequences were absent from training. After excluding 10 reused host species identifiers, full-model AUROC was 0.625 on 663 pairs. After excluding pairs with either entity at or below the prespecified sequence-distance threshold, AUROC was 0.644 and AUPRC 0.427 on 805 pairs (35.2% positive). Across eligible groups, median per-phage and per-host AUROCs were 0.625 and 0.663, respectively, but ranges were wide. These secondary point estimates are compatible with possible ranking signal, but they do not convert the inconclusive primary evaluation into confirmatory evidence of cross-source transfer.",
    )
    replace_paragraph(
        doc,
        "The central result now has two parts.",
        "The central result has two parts. Within sequence-covered NCBI_HR, discrimination decreased as taxonomic and sequence-cluster separation became stricter. Across the NCBI_HR-to-StaphStudy boundary, the frozen full model produced an AUROC point estimate of 0.637, but its primary cluster-aware interval included 0.5 and calibration deteriorated sharply. Cross-source discrimination is therefore unresolved rather than demonstrated. The similar sequence-only point estimate indicates that the four VHIP-supplied columns did not solely create the observed ranking pattern, but this sensitivity analysis does not rescue the inconclusive primary result. Together, the analyses establish that leakage control changes apparent within-source difficulty while leaving transportability uncertain.",
    )
    replace_paragraph(
        doc,
        "Composition-based WIsH",
        "Composition-based WIsH, exact-word PHIST, protein-centric RaFAH and CHERRY, and network-integrated predictors capture different signals (Ahlgren et al., 2017; Galiez et al., 2017; Coutinho et al., 2021; Shang and Sun, 2022; Wang et al., 2020; Zielezinski et al., 2022). A 27-tool benchmark likewise found strong context dependence across datasets and task formulations (Shang et al., 2025). The present results reinforce that warning: source prevalence ranged widely, the four-feature transfer audit was weak, and the full sequence model's point estimate exceeded 0.5 without excluding chance under cluster-aware uncertainty, while its probabilities were poorly calibrated. The study therefore cannot claim reliable cross-source ordering, clinically usable probabilities, or transferable thresholds.",
    )
    replace_paragraph(
        doc,
        "The GNN result remains within-source.",
        "The GNN result remains within-source. Message passing modestly improved the otherwise identical no-graph neural model within NCBI_HR, but the GNN remained below the fixed GBM in every regime and was not frozen for the StaphStudy test. The cross-source evaluation tests only the fixed tree pipeline and does not establish its transfer; it provides no validation of the model-class ranking, PHIST comparison, RaFAH-inspired proxy, or downstream set-cover selections. Those comparisons should not be generalized beyond their aligned NCBI_HR test rows.",
    )
    replace_paragraph(
        doc,
        "The downstream analyses are transformations",
        "The downstream analyses are transformations of NCBI_HR model output, not independent evidence of generalization or mechanism. Neither the above-chance cross-source point estimate nor the inconclusive primary interval validates the 176-phage set cover, its thresholds, or the deterministic resistance equations. Sparse untested cells, taxon-level targets, very large selections, and poor probability transport prevent a therapeutic-cocktail claim. The equations impose rather than learn resistance and kinetic parameters; their universal rebound within the sampled grid cannot estimate rebound probability, validate a mechanism, compare therapies, or strengthen the introductory AMR rationale. Biological claims would require measured adsorption, killing, mutation, fitness, cross-resistance, and prospective time courses.",
    )
    replace_paragraph(
        doc,
        "PrecisionPhage provides an auditable",
        "PrecisionPhage provides an auditable within-NCBI_HR leakage hierarchy and a frozen, retraining-free NCBI_HR-to-StaphStudy sequence evaluation. The cross-source AUROC point estimate was above 0.5, but the primary cluster-aware interval included chance and calibration was poor; cross-source discrimination therefore remains inconclusive. The analysis did not validate downstream set cover, resistance biology, or therapeutic efficacy. Independent prospective sequence-covered matrices and empirical experiments remain necessary before predictive deployment or biological claims.",
    )
    replace_paragraph(
        doc,
        "Table 8. Frozen NCBI_HR-to-StaphStudy evaluation.",
        "Table 8. Frozen NCBI_HR-to-StaphStudy evaluation. Confidence intervals are two-way sequence-cluster bootstrap intervals (2,000 valid replicates); the primary full-model AUROC interval includes 0.5. Sequence-only excludes the four VHIP-supplied columns; the four-feature row uses only those columns. Sensitivity rows are contextual and do not supersede the primary result.",
    )

    paragraphs = doc.paragraphs
    intro_index = next(i for i, p in enumerate(paragraphs) if p.text == "1 Introduction")
    declarations_index = next(i for i, p in enumerate(paragraphs)
                              if p.text == "Conflict of Interest")
    word_count = sum(
        len(re.findall(r"\b[\w–-]+\b", p.text))
        for p in paragraphs[intro_index:declarations_index]
        if not p.style.name.startswith("Heading")
    )
    replace_paragraph(
        doc,
        "Article type: Original Research.",
        f"Article type: Original Research. Manuscript statistics: approximately {word_count:,} main-text words; 5 figures; 8 tables.",
    )

    # Keep the reference list compact enough to avoid a nearly empty spill page.
    paragraphs = doc.paragraphs
    references_index = next(i for i, p in enumerate(paragraphs)
                            if p.text == "References")
    tables_index = next(i for i, p in enumerate(paragraphs)
                        if p.text == "Tables")
    for paragraph in paragraphs[references_index + 1:tables_index]:
        if paragraph.text.strip():
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(6)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
