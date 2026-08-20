#!/usr/bin/env python3
"""Integrate the locked fixed-bank validation into the latest manuscript."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent.parent
SOURCE = (ROOT / "submission" /
          "PrecisionPhage_Frontiers_Original_Research_cross_source_uncertainty_reframed.docx")
OUTPUT = (ROOT / "submission" /
          "PrecisionPhage_Frontiers_Original_Research_external_cocktail_validated.docx")


def set_text(paragraph: Paragraph, text: str) -> None:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def insert_before(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new = paragraph.insert_paragraph_before(text)
    new.style = style
    # The Frontiers template's Heading styles carry automatic list formatting.
    # Existing manuscript headings explicitly suppress it with numId=0.
    if style.startswith("Heading"):
        p_pr = new._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "0")
        num_pr.append(num_id)
        p_pr.append(num_pr)
    return new


def find(doc: Document, exact: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise RuntimeError(f"Paragraph not found: {exact}")


def find_startswith(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Paragraph prefix not found: {prefix}")


def mark_header_row(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_repeatable_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    mark_header_row(table)


def main() -> None:
    doc = Document(SOURCE)
    set_text(doc.paragraphs[0],
             "Leakage-controlled phage-host prediction supports fixed-bank cocktail prioritization for unseen Escherichia hosts")

    abstract = find(doc, "Abstract")._p.getnext()
    abstract_paragraph = Paragraph(abstract, find(doc, "Abstract")._parent)
    set_text(abstract_paragraph,
             "Computational prediction of bacteriophage-host interactions may prioritize experimental host-range testing, but related genomes and source structure can inflate evaluation. We present PrecisionPhage, a reproducible leakage-controlled framework using 8,849 experimentally assayed pairs from three study labels. Within the 1,947-pair sequence-covered NCBI_HR subset, pooled fixed-tree AUROC decreased from 0.959 for eligible unseen host species to 0.785 across five dual cold-start blocks. Retraining-free evaluation on all 1,053 sequence-covered StaphStudy pairs yielded AUROC 0.637 (two-way sequence-cluster bootstrap 95% CI 0.498-0.739), AUPRC 0.391 (0.201-0.591), and expected calibration error 0.499. Because the primary AUROC interval includes 0.5 and calibration was poor, cross-source discrimination remains unresolved and transportable probabilities were not established. Separately, under a protocol locked before external outcome access, we trained on the dense 402-host by 96-phage Picard matrix and froze scores for 100 unseen ColoColi hosts. The prespecified score ranked productive activity of the externally selected, experimentally tested published cocktails with AUROC 0.807 (host-cluster bootstrap 95% CI 0.693-0.906) and average precision 0.951 (0.904-0.983). This result validates only candidate-cocktail ranking within a fixed Escherichia phage bank; it does not validate de novo PrecisionPhage cocktail selection. Together, the analyses establish a leakage-controlled within-source difficulty hierarchy, unresolved cross-source discrimination for a sparse curated source, and bounded fixed-bank prioritization. They do not establish general cross-database transport, portable calibrated probabilities, unseen-phage transfer, clinical efficacy, or resistance suppression.")

    set_text(find(doc, "Here we separate questions requiring different evidence. First, a leave-one-study-out audit tests transfer across three source labels using the four features present for every pair. Second, within sequence-covered NCBI_HR, taxonomic and sequence-cluster holdouts compare a fixed gradient-boosted tree with GraphSAGE, PHIST, and a RaFAH-inspired proxy. Third, a protocol locked before prediction evaluates the unchanged full sequence pipeline retraining-free on sequence-covered StaphStudy pairs. Out-of-fold NCBI_HR scores are additionally examined in retrospective set cover. A deterministic resistance grid is only an ancillary illustration of imposed assumptions. This computational Original Research reports positive and negative validation results but does not claim wet-lab, prospective, or clinical validation."),
             "Here we separate four questions and their claim boundaries. First, a leave-one-study-out audit uses the four features present for every pair; it measures source heterogeneity but does not validate the full 24-feature sequence model. Second, within sequence-covered NCBI_HR, taxonomic and sequence-cluster holdouts compare a fixed gradient-boosted tree with GraphSAGE, PHIST, and a RaFAH-inspired proxy; these comparisons remain within one source. Third, a protocol locked before prediction evaluates the unchanged full sequence pipeline retraining-free on sequence-covered StaphStudy pairs. Its primary interval includes chance and its calibration is poor, so it is retained as inconclusive rather than evidence of general cross-source prediction. Fourth, a separate protocol locked before external outcome access trains on the dense Picard Escherichia matrix and scores published cocktails on 100 unseen ColoColi hosts with experimental outcomes (Gaborieau et al., 2024). Because the external study selected those compositions, this analysis validates only fixed-bank candidate ranking, not de novo PrecisionPhage cocktail selection. Out-of-fold NCBI_HR scores are additionally examined in retrospective set cover, and a deterministic resistance grid is only an ancillary illustration of imposed assumptions. This computational Original Research does not claim broad cross-database generalization, transportable calibrated probabilities, prospective clinical validation, or therapeutic efficacy.")

    # Insert the independent validation method and renumber following subsections.
    models_heading = find(doc, "2.5 Models")
    insert_before(models_heading, "2.5 Locked external fixed-bank cocktail validation", "Heading 2")
    insert_before(models_heading,
                  "Before reading the held-out cocktail outcome and composition tables, we locked a second protocol using the independent Escherichia resource of Gaborieau et al. (2024). Development used 38,531 observed interactions from the 402-host by 96-phage Picard matrix; the test comprised 100 previously unseen pathogenic ColoColi hosts. For each phage, a regularized one-hot logistic model used five host covariates shared across collections (phylogroup, sequence type, LPS type, O type, and H type), and a distance-weighted nearest-neighbour model used core-genome distance to Picard hosts. Closely related development hosts were kept in the same cross-validation group by connected components at the published 10^-4 distance threshold. Logistic regularization and neighbour count were selected globally by mean phage-level average precision using Picard labels only. After refitting on all development hosts, the two components were rank-normalized across test hosts and averaged, producing 9,600 frozen pair scores. The prespecified cocktail score was 1 - product(1 - p_i) over three constituents; it was used only as a ranking score and did not assert biological independence. Primary outcomes were AUROC and average precision for productive cocktail activity (minimum lytic concentration greater than zero), with 2,000 host-cluster bootstrap replicates. The external study selected and experimentally tested the cocktail compositions; PrecisionPhage did not select those tested compositions.")
    for old, new in [("2.5 Models", "2.6 Models"),
                     ("2.6 External baselines", "2.7 External baselines"),
                     ("2.7 Statistical analysis", "2.8 Statistical analysis"),
                     ("2.8 Set-cover prioritization and ancillary deterministic illustration", "2.9 Set-cover prioritization and ancillary deterministic illustration"),
                     ("2.9 Set-cover formulation and illustrative scenario equations", "2.10 Set-cover formulation and illustrative scenario equations"),
                     ("2.10 Software, reproducibility, and AI assistance", "2.11 Software, reproducibility, and AI assistance")]:
        set_text(find(doc, old), new)

    # Insert the external result before the existing GBM/GNN subsection.
    target = find(doc, "3.3 Within NCBI_HR, the gradient-boosted tree has higher pooled AUROC than the graph neural network")
    insert_before(target, "3.3 Locked fixed-bank scores rank productive candidate cocktails on unseen Escherichia hosts", "Heading 2")
    insert_before(target,
                  "Among the 100 experimentally tested published cocktails, 82 produced lytic activity. The frozen PrecisionPhage cocktail score yielded AUROC 0.807 (host-cluster bootstrap 95% CI 0.693-0.906), average precision 0.951 (0.904-0.983), Brier score 0.127, and expected calibration error 0.079 (Table 9). The maximum-constituent sensitivity score gave AUROC 0.800. In a secondary analysis restricted to the 300 phage-host constituents chosen by the published recommender, ensemble AUROC was 0.668 (two-way host-cluster and phage-genus bootstrap 95% CI 0.559-0.775) and average precision was 0.737 (0.599-0.870). These estimates support ranking of externally defined candidate cocktails for unseen hosts within this fixed Escherichia phage bank. They do not establish calibration outside this task, transfer to a different bank or organism, or de novo cocktail selection. Because Gaborieau et al. selected the tested compositions, the result is not evidence of therapeutic efficacy.")
    for old, new in [("3.3 Within NCBI_HR, the gradient-boosted tree has higher pooled AUROC than the graph neural network", "3.4 Within NCBI_HR, the gradient-boosted tree has higher pooled AUROC than the graph neural network"),
                     ("3.4 Genome composition and nucleotide homology dominate", "3.5 Genome composition and nucleotide homology dominate"),
                     ("3.5 Within NCBI_HR, PrecisionPhage has higher pooled AUROC than the evaluated baselines", "3.6 Within NCBI_HR, PrecisionPhage has higher pooled AUROC than the evaluated baselines"),
                     ("3.6 Exact optimization of predicted set cover produces large selections", "3.7 Exact optimization of predicted set cover produces large selections"),
                     ("3.7 Prespecified equations yield conditional rebound across the illustrative grid", "3.8 Prespecified equations yield conditional rebound across the illustrative grid")]:
        set_text(find(doc, old), new)

    discussion_second = find(doc, "Composition-based WIsH, exact-word PHIST, protein-centric RaFAH and CHERRY, and network-integrated predictors capture different signals (Ahlgren et al., 2017; Galiez et al., 2017; Coutinho et al., 2021; Shang and Sun, 2022; Wang et al., 2020; Zielezinski et al., 2022). A 27-tool benchmark likewise found strong context dependence across datasets and task formulations (Shang et al., 2025). The present results reinforce that warning: source prevalence ranged widely, the four-feature transfer audit was weak, and the full sequence model's point estimate exceeded 0.5 without excluding chance under cluster-aware uncertainty, while its probabilities were poorly calibrated. The study therefore cannot claim reliable cross-source ordering, clinically usable probabilities, or transferable thresholds.")
    set_text(discussion_second,
             "Composition-based WIsH, exact-word PHIST, protein-centric RaFAH and CHERRY, and network-integrated predictors capture different signals (Ahlgren et al., 2017; Galiez et al., 2017; Coutinho et al., 2021; Shang and Sun, 2022; Wang et al., 2020; Zielezinski et al., 2022). A 27-tool benchmark likewise found strong context dependence across datasets and task formulations (Shang et al., 2025). The present results reinforce that warning: source prevalence ranged widely, the four-feature cross-source audit produced weak performance, and the full sequence model's point estimate exceeded 0.5 without excluding chance under cluster-aware uncertainty, while its probabilities were poorly calibrated. The study therefore cannot claim reliable cross-source ordering, clinically usable probabilities, or transferable thresholds.")
    insert_before(discussion_second,
                  "The positive Picard-to-ColoColi result narrows rather than removes the generalization problem. Unlike NCBI_HR, where the median phage appears in one observed pair and the sequence-covered set is 76.4% positive, Picard is an almost complete fixed-bank matrix with 20.8% positives. Repeated observations per phage and host make within-bank susceptibility learnable, and the locked result shows that the resulting scores rank experimentally productive candidate cocktails for unseen hosts. This is directly relevant to laboratory triage, but its boundary is important: the phage bank and organismal genus were fixed, and the external study—not PrecisionPhage—selected the tested compositions. Thus, the result supports cocktail prioritization for testing, not autonomous therapeutic design (Gaborieau et al., 2024).")
    set_text(find(doc, "The downstream analyses are transformations of NCBI_HR model output, not independent evidence of generalization or mechanism. Neither the above-chance cross-source point estimate nor the inconclusive primary interval validates the 176-phage set cover, its thresholds, or the deterministic resistance equations. Sparse untested cells, taxon-level targets, very large selections, and poor probability transport prevent a therapeutic-cocktail claim. The equations impose rather than learn resistance and kinetic parameters; their universal rebound within the sampled grid cannot estimate rebound probability, validate a mechanism, compare therapies, or strengthen the introductory AMR rationale. Biological claims would require measured adsorption, killing, mutation, fitness, cross-resistance, and prospective time courses."),
             "The NCBI_HR downstream analyses remain transformations of sparse within-source model output and are not independent evidence of mechanism. The independent Escherichia score does not validate the 176-phage NCBI_HR set cover, its thresholds, or the deterministic resistance equations. Sparse untested NCBI_HR cells, taxon-level targets, and very large selections prevent a therapeutic formulation claim. The equations impose rather than learn resistance and kinetic parameters; their universal rebound within the sampled grid cannot estimate rebound probability, validate a mechanism, compare therapies, or strengthen the introductory AMR rationale. Biological claims would require measured adsorption, killing, mutation, fitness, cross-resistance, and prospective time courses.")
    set_text(find(doc, "The frozen StaphStudy test addresses the previous absence of cross-source sequence evaluation, but it does not eliminate the generalization limitation. StaphStudy was already part of VHRnet and the earlier four-feature audit, so the test was source-held-out rather than prospectively blinded or independently collected. It contains 39 phages, 27 host reference genomes, and only 15 phage sequence clusters; 10 host species identifiers recurred in training, and four test genomes on each axis fell within the prespecified similarity threshold, although the strict 805-pair subset gave a similar result. The full-model AUROC interval included chance and calibration was poor. Further limitations include evaluation of only eligible internal groups, five cold-start blocks, repeated entities, no wet-lab or prospective validation, a RaFAH-inspired proxy rather than published RaFAH, sparse cocktail assays, and unfitted resistance parameters. The work supports a leakage and transportability analysis, not a deployable predictor, resistance model, or efficacy claim."),
             "The frozen StaphStudy test addresses the previous absence of cross-source sequence evaluation, but it does not establish general Staphylococcus transfer. StaphStudy was already part of VHRnet and the earlier four-feature audit, so the test was source-held-out rather than prospectively blinded or independently collected. It contains 39 phages, 27 host reference genomes, and only 15 phage sequence clusters; 10 host species identifiers recurred in training, and four test genomes on each axis fell within the prespecified similarity threshold, although the strict 805-pair subset gave a similar result. The full-model AUROC interval included chance and calibration was poor. Further limitations include evaluation of only eligible internal groups, five cold-start blocks, repeated entities, no prospective PrecisionPhage-selected cocktail experiment, a RaFAH-inspired proxy rather than published RaFAH, sparse NCBI_HR cocktail assays, and unfitted resistance parameters. The work supports leakage-aware prediction and bounded fixed-bank prioritization, not a deployable clinical predictor, resistance model, or efficacy claim.")
    set_text(find(doc, "PrecisionPhage provides an auditable within-NCBI_HR leakage hierarchy and a frozen, retraining-free NCBI_HR-to-StaphStudy sequence evaluation. The cross-source AUROC point estimate was above 0.5, but the primary cluster-aware interval included chance and calibration was poor; cross-source discrimination therefore remains inconclusive. The analysis did not validate downstream set cover, resistance biology, or therapeutic efficacy. Independent prospective sequence-covered matrices and empirical experiments remain necessary before predictive deployment or biological claims."),
             "PrecisionPhage provides an auditable within-NCBI_HR leakage hierarchy and an inconclusive frozen NCBI_HR-to-StaphStudy test. Because the primary cross-source interval included chance and calibration was poor, the study does not establish broad cross-database discrimination, portable calibrated probabilities, or transferable decision thresholds. In a separate bounded task, the locked fixed-bank score ranked externally selected, experimentally tested candidate cocktails on unseen Escherichia hosts with AUROC 0.807 (95% CI 0.693-0.906). This supports candidate prioritization for laboratory testing within that bank, not de novo PrecisionPhage cocktail design or therapeutic efficacy. Prospective testing of model-selected cocktails and measured resistance remain necessary before clinical or efficacy claims.")
    set_text(find_startswith(doc, "The direct interaction table is available with the VHIP publication"),
             "The direct VHRnet interaction table is available with the VHIP publication (Bastien et al., 2024), and the StaphStudy genomes are available through the VHIP example bundle and GenBank accessions MZ417315-MZ417354 reported by Göller et al. (2021). The independent Picard matrix, genomes, metadata, and ColoColi cocktail outcomes are available through the GitHub and Zenodo resources of Gaborieau et al. (2024). Frozen analysis code, both dated external-validation protocols, predictions, manifests, model artifacts, overlap audits, tables, figures, checksums, and reproduction instructions are available at https://github.com/JaivardhanSamadhiya/frontiers_submission. Large genome FASTA files and external repository checkouts are excluded because of size; pinned commits and staging identifiers are recorded.")

    # Insert the new reference alphabetically before Galiez.
    galiez = next(p for p in doc.paragraphs if p.text.startswith("Galiez, C."))
    insert_before(galiez,
                  "Gaborieau, B., Vaysset, H., Tesson, F., Charachon, I., Dib, N., Bernier, J., et al. (2024). Prediction of strain level phage-host interactions across the Escherichia genus using only genomic information. Nat. Microbiol. 9, 2847-2861. doi: 10.1038/s41564-024-01832-5")

    # Caption and table are inserted immediately before Figure Legends.
    figure_legends = find(doc, "Figure Legends")
    caption = insert_before(figure_legends,
                            "Table 9. Locked Picard-to-ColoColi fixed-bank evaluation. The primary row scores 100 experimentally tested published cocktails on unseen hosts; confidence intervals use 2,000 host-cluster bootstrap replicates. The constituent row is secondary and uses two-way host-cluster and phage-genus bootstrap. Gaborieau et al. selected the tested compositions.")
    # Start the compact final table on a new page so its caption cannot be
    # orphaned by Word's pagination engine.
    caption.paragraph_format.page_break_before = True
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Evaluation", "N", "Positive", "AUROC (95% CI)",
               "Average precision (95% CI)", "Status"]
    data = [
        ["Cocktail score", "100", "82", "0.807 (0.693-0.906)",
         "0.951 (0.904-0.983)", "Primary"],
        ["Selected constituent pairs", "300", "170", "0.668 (0.559-0.775)",
         "0.737 (0.599-0.870)", "Secondary"],
    ]
    for j, value in enumerate(headers):
        table.rows[0].cells[j].text = value
        for run in table.rows[0].cells[j].paragraphs[0].runs:
            run.bold = True
    for values in data:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = value
    set_repeatable_table_geometry(table, [1.20, 0.45, 0.60, 1.35, 1.55, 0.70])
    caption._p.addnext(table._tbl)

    # Update manuscript statistics from the final main-text section.
    paragraphs = [p.text for p in doc.paragraphs]
    intro = next(i for i, text in enumerate(paragraphs) if text.strip() == "1 Introduction")
    conflict = next(i for i, text in enumerate(paragraphs) if text.strip() == "Conflict of Interest")
    main_words = len(re.findall(r"\b[\w'-]+\b", " ".join(paragraphs[intro:conflict])))
    stats = next(p for p in doc.paragraphs if p.text.startswith("Article type:"))
    set_text(stats, f"Article type: Original Research. Manuscript statistics: approximately {main_words:,} main-text words; 5 figures; 9 tables.")
    set_text(doc.paragraphs[5], "Running title: Bounded fixed-bank cocktail score validation")

    set_text(find(doc, "We first evaluated source transfer with the full-table four-feature audit, then evaluated stricter entity holdouts within sequence-covered NCBI_HR. The newly locked test applied the full sequence representation from NCBI_HR to StaphStudy without retraining. Model-class, comparator, set-cover, and deterministic-scenario analyses otherwise remain based on NCBI_HR and are interpreted accordingly."),
             "We first evaluated source heterogeneity with the full-table four-feature audit, then evaluated stricter entity holdouts within sequence-covered NCBI_HR. The frozen sequence test applied the full NCBI_HR representation to StaphStudy without retraining; its discrimination and calibration are reported as an unresolved cross-source result. The separate fixed-bank test evaluated only ranking of externally defined candidate cocktails from Picard development data on unseen ColoColi hosts. Model-class, comparator, set-cover, and deterministic-scenario analyses otherwise remain based on NCBI_HR and are interpreted accordingly.")

    set_text(find(doc, "3.1 Cross-source transfer is limited, and within-source performance declines under stricter holdouts"),
             "3.1 Four-feature cross-source performance is weak, and within-source performance declines under stricter holdouts")

    set_text(find(doc, "Before generating held-out predictions, we locked the test cohort, feature schema, fixed XGBoost configuration, exclusion rules, and metrics in a dated protocol. A StandardScaler and the fixed 400-tree model were fitted once on all 1,947 NCBI_HR rows without access to StaphStudy labels, then applied without retraining or threshold optimization to all 1,053 StaphStudy rows. Primary outcomes were pooled AUROC, AUPRC, Brier score, and 10-bin expected calibration error. Ninety-five-percent intervals used 2,000 valid two-way bootstrap replicates that independently resampled phage and host sequence clusters; row-wise resampling was not used. Secondary analyses excluded the 10 test host species identifiers seen in training (663 pairs) and excluded any pair for which either genome fell at or below the prespecified 0.05 Mash-form distance to a training genome (805 pairs). Exact-sequence overlap was zero on both axes; four test phages and four test hosts fell at or below the distance threshold. Because StaphStudy was already a component of VHRnet and its labels had appeared in the earlier four-feature audit, this evaluation is source-held-out but not prospectively blinded or independently collected."),
             "Before generating held-out predictions, we locked the test cohort, feature schema, fixed XGBoost configuration, exclusion rules, and metrics in a dated protocol. A StandardScaler and the fixed 400-tree model were fitted once on all 1,947 NCBI_HR rows without access to StaphStudy labels, then applied without retraining or threshold optimization to all 1,053 StaphStudy rows. Primary outcomes were pooled AUROC, AUPRC, Brier score, and 10-bin expected calibration error. Ninety-five-percent intervals used 2,000 valid two-way bootstrap replicates that independently resampled phage and host sequence clusters; row-wise resampling was not used. Secondary analyses excluded the 10 test host species identifiers seen in training (663 pairs) and excluded any pair for which either genome fell at or below the prespecified 0.05 Mash-form distance to a training genome (805 pairs). Exact-sequence overlap was zero on both axes; four test phages and four test hosts fell at or below the distance threshold. Because StaphStudy was already a component of VHRnet and its labels had appeared in the earlier four-feature audit, this evaluation is source-held-out but not prospectively blinded or independently collected.")

    interaction_data = find_startswith(doc, "We used the 8,849 laboratory-tested")
    set_text(
        interaction_data,
        interaction_data.text.replace(
            "distributed with the Virus-host interactions predictor (VHIP) study",
            "in the VHRnet interaction table distributed with the Virus-host interactions predictor (VHIP) study",
        ),
    )

    paragraphs = [p.text for p in doc.paragraphs]
    intro = next(i for i, text in enumerate(paragraphs) if text.strip() == "1 Introduction")
    conflict = next(i for i, text in enumerate(paragraphs) if text.strip() == "Conflict of Interest")
    main_words = len(re.findall(r"\b[\w'-]+\b", " ".join(paragraphs[intro:conflict])))
    set_text(stats, f"Article type: Original Research. Manuscript statistics: approximately {main_words:,} main-text words; 5 figures; 9 tables.")

    # The Frontiers template implements its provisional footer with anchored
    # drawing text boxes. Give those drawing containers descriptive metadata so
    # automated accessibility checks do not treat them as unlabeled images.
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer,
                       section.even_page_footer):
            for drawing_properties in footer._element.xpath(".//wp:docPr"):
                if not (drawing_properties.get("descr") or
                        drawing_properties.get("title")):
                    drawing_properties.set(
                        "descr", "Frontiers provisional manuscript footer")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"main_text_words={main_words}; references=25; tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
